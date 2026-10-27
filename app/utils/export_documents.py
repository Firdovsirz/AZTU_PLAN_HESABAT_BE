"""Formal PDF / Excel rendering for plan & report exports.

Rendering only: app.crud.export collects and scopes the data, then hands an
ExportDocument to render_pdf / render_xlsx. Nothing here touches the database.

Row dicts carry these keys (missing values are None):
  plans:   unit, full_name, fin_kod, serial, work_year, row_number, activities,
           work_desc, goal, deadline, created_at
  reports: unit, full_name, fin_kod, serial, work_year, activities,
           result_indicator, done_percentage, self_score_label, admin_score,
           ai_score, duration, submitted_at, status, note, has_document
"""
import io
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from itertools import groupby
from pathlib import Path
from typing import Callable
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (
    CondPageBreak,
    HRFlowable,
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets"
LOGO_PATH = ASSETS_DIR / "aztu_logo.png"
FONT_DIR = ASSETS_DIR / "fonts"

# Azerbaijan has no DST, so a fixed offset is exact and needs no tzdata.
BAKU_TZ = timezone(timedelta(hours=4))

UNIVERSITY_NAME = "Azərbaycan Texniki Universiteti"
# Spelled out because str.upper() maps "i" to "I", not the Azerbaijani "İ".
UNIVERSITY_NAME_UPPER = "AZƏRBAYCAN TEXNİKİ UNİVERSİTETİ"
SYSTEM_NAME = "Plan-Hesabat İnformasiya Sistemi"

# Palette: the navy is sampled from the AzTU logo.
NAVY_HEX = "1B2C5D"
NAVY = colors.HexColor(f"#{NAVY_HEX}")
INK = colors.HexColor("#1F2937")
MUTED = colors.HexColor("#6B7280")
BORDER_HEX = "C9CFDD"
BORDER = colors.HexColor(f"#{BORDER_HEX}")
ZEBRA_HEX = "F3F5FA"
ZEBRA = colors.HexColor(f"#{ZEBRA_HEX}")
TINT = colors.HexColor("#EEF1F8")

DURATION_LABELS = {0: "Vaxtından gec", 1: "Vaxtında", 2: "Vaxtından əvvəl"}
STATUS_LABELS = {"done": "Tamamlanıb", "submitted": "Təhvil verilib", "pending": "İcradadır"}
EMPTY_TEXT = "Seçilmiş meyarlara uyğun məlumat tapılmadı."


@dataclass
class ExportDocument:
    title: str
    # Key/value lines shown under the title (scope, person, year, ...).
    details: list[tuple[str, str]]
    # None means the section was not requested; [] means requested but empty.
    plans: list[dict] | None = None
    reports: list[dict] | None = None
    # Several people → show name columns; several units → group/label by unit.
    multi_person: bool = False
    multi_unit: bool = False
    signer_label: str = "Tərtib etdi"
    signer_name: str | None = None
    generated_at: datetime = field(default_factory=lambda: datetime.now(BAKU_TZ))


# ---------------------------------------------------------------------------
# Column specs shared by both formats. pdf_width is a relative weight (0 keeps
# the column out of the PDF); xlsx_width is an Excel character width (0 keeps
# it out of the workbook).
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class _Col:
    header: str
    value: Callable[[dict], object]
    pdf_width: float = 0
    xlsx_width: float = 0
    person: bool = False
    unit: bool = False
    center: bool = False


def _get(key: str) -> Callable[[dict], object]:
    return lambda row: row.get(key)


_UNIT_COL = _Col("Struktur bölmə", _get("unit"), xlsx_width=34, unit=True)
_NAME_COL = _Col("Soyadı, adı, ata adı", _get("full_name"), pdf_width=1.45, xlsx_width=32, person=True)
_FIN_COL = _Col("FİN", _get("fin_kod"), xlsx_width=11, person=True, center=True)

PLAN_COLUMNS = [
    _UNIT_COL,
    _NAME_COL,
    _FIN_COL,
    _Col("Plan №", _get("serial"), pdf_width=1.25, xlsx_width=18),
    _Col("İl", _get("work_year"), pdf_width=0.45, xlsx_width=7, center=True),
    _Col("Sıra", _get("row_number"), xlsx_width=7, center=True),
    _Col("Fəaliyyət növləri", _get("activities"), pdf_width=1.8, xlsx_width=34),
    _Col("İşin qısa təsviri", _get("work_desc"), pdf_width=2.6, xlsx_width=50),
    _Col("Hədəf", _get("goal"), pdf_width=2.0, xlsx_width=40),
    _Col("Son tarix", _get("deadline"), pdf_width=0.8, xlsx_width=12, center=True),
    _Col("Yaradılma tarixi", _get("created_at"), xlsx_width=14, center=True),
]

REPORT_COLUMNS = [
    _UNIT_COL,
    _NAME_COL,
    _FIN_COL,
    _Col("Plan №", _get("serial"), pdf_width=1.25, xlsx_width=18),
    _Col("İl", _get("work_year"), xlsx_width=7, center=True),
    _Col("Fəaliyyət növləri", _get("activities"), pdf_width=1.45, xlsx_width=34),
    _Col("Nəticə indikatoru", _get("result_indicator"), pdf_width=2.0, xlsx_width=44),
    _Col("Uğur faizi", _get("done_percentage"), pdf_width=0.8, xlsx_width=14, center=True),
    _Col("İcraçının qiyməti", _get("self_score_label"), pdf_width=0.9, xlsx_width=22),
    _Col("Keyfiyyət şöbəsinin qiyməti (0–100)", _get("admin_score"), pdf_width=0.8, xlsx_width=16, center=True),
    _Col("Süni intellekt qiyməti", _get("ai_score"), xlsx_width=14, center=True),
    _Col("İcra müddəti", lambda r: DURATION_LABELS.get(r.get("duration")), pdf_width=0.85, xlsx_width=16, center=True),
    _Col("Təhvil tarixi", _get("submitted_at"), pdf_width=0.8, xlsx_width=14, center=True),
    _Col("Status", lambda r: STATUS_LABELS.get(r.get("status")), pdf_width=0.85, xlsx_width=16, center=True),
    _Col("Qeyd", _get("note"), xlsx_width=30),
    _Col("Sənəd", lambda r: "Əlavə edilib" if r.get("has_document") else None, xlsx_width=13, center=True),
]


def _visible(columns: list[_Col], doc: ExportDocument, fmt: str) -> list[_Col]:
    out = []
    for col in columns:
        if (col.pdf_width if fmt == "pdf" else col.xlsx_width) <= 0:
            continue
        if col.person and not doc.multi_person:
            continue
        if col.unit and not doc.multi_unit:
            continue
        out.append(col)
    return out


def _fmt_date(value) -> str:
    return value.strftime("%d.%m.%Y") if isinstance(value, (date, datetime)) else ""


def _summary(doc: ExportDocument) -> list[tuple[str, str]]:
    base = doc.plans if doc.plans is not None else (doc.reports or [])
    items = [("Planların sayı", str(len(base)))]
    if doc.reports is not None:
        submitted = sum(1 for r in doc.reports if r.get("status") in ("submitted", "done"))
        done = sum(1 for r in doc.reports if r.get("status") == "done")
        scores = [r["admin_score"] for r in doc.reports if r.get("admin_score") is not None]
        items += [
            ("Təhvil verilmiş hesabatlar", str(submitted)),
            ("Tamamlanmış hesabatlar", str(done)),
            ("Orta keyfiyyət qiyməti", f"{sum(scores) / len(scores):.1f}" if scores else "—"),
        ]
    return items


def _sections(doc: ExportDocument):
    """(numeral, heading, sheet name, columns, rows, noun) per requested section."""
    numerals = iter(["I", "II"])
    if doc.plans is not None:
        yield next(numerals), "İş planları", "Planlar", PLAN_COLUMNS, doc.plans, "plan"
    if doc.reports is not None:
        yield next(numerals), "Hesabatlar", "Hesabatlar", REPORT_COLUMNS, doc.reports, "hesabat"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_FONT = "DejaVuSansCondensed"
_FONT_BOLD = "DejaVuSansCondensed-Bold"
_PDF_MARGIN = 14 * mm
# A single table cell longer than this is cut in the PDF (the Excel file
# always carries the full text) so one runaway field can't fill many pages.
_PDF_MAX_CELL_CHARS = 2500
_LOGO_W, _LOGO_H = ImageReader(str(LOGO_PATH)).getSize()


def _register_fonts() -> None:
    # The built-in PDF fonts lack Azerbaijani glyphs (ə, ğ, ı, ş ...), so a
    # bundled Unicode TTF is embedded instead.
    if _FONT in pdfmetrics.getRegisteredFontNames():
        return
    pdfmetrics.registerFont(TTFont(_FONT, str(FONT_DIR / "DejaVuSansCondensed.ttf")))
    pdfmetrics.registerFont(TTFont(_FONT_BOLD, str(FONT_DIR / "DejaVuSansCondensed-Bold.ttf")))
    pdfmetrics.registerFontFamily(_FONT, normal=_FONT, bold=_FONT_BOLD, italic=_FONT, boldItalic=_FONT_BOLD)


def _styles() -> dict[str, ParagraphStyle]:
    base = ParagraphStyle("base", fontName=_FONT, fontSize=8, leading=10, textColor=INK)
    return {
        "kicker": ParagraphStyle("kicker", base, fontName=_FONT_BOLD, fontSize=9, leading=12, textColor=NAVY, alignment=TA_RIGHT),
        "title": ParagraphStyle("title", base, fontName=_FONT_BOLD, fontSize=17, leading=21, textColor=NAVY, alignment=TA_RIGHT),
        "system": ParagraphStyle("system", base, fontSize=8.5, leading=11, textColor=MUTED, alignment=TA_RIGHT),
        "label": ParagraphStyle("label", base, fontSize=8, textColor=MUTED),
        "value": ParagraphStyle("value", base, fontName=_FONT_BOLD, fontSize=8.5, leading=11),
        "stat_value": ParagraphStyle("stat_value", base, fontName=_FONT_BOLD, fontSize=14, leading=17, textColor=NAVY, alignment=TA_CENTER),
        "stat_label": ParagraphStyle("stat_label", base, fontSize=7.5, leading=9, textColor=MUTED, alignment=TA_CENTER),
        "h2": ParagraphStyle("h2", base, fontName=_FONT_BOLD, fontSize=12, leading=15, textColor=NAVY, spaceBefore=12, spaceAfter=6),
        "h3": ParagraphStyle("h3", base, fontName=_FONT_BOLD, fontSize=9.5, leading=12, textColor=INK, spaceBefore=8, spaceAfter=4),
        "th": ParagraphStyle("th", base, fontName=_FONT_BOLD, fontSize=7.5, leading=9, textColor=colors.white, alignment=TA_CENTER),
        "td": ParagraphStyle("td", base, fontSize=7.8, leading=9.6, alignment=TA_LEFT),
        "td_center": ParagraphStyle("td_center", base, fontSize=7.8, leading=9.6, alignment=TA_CENTER),
        "empty": ParagraphStyle("empty", base, fontSize=8.5, textColor=MUTED, spaceBefore=2, spaceAfter=6),
        "note": ParagraphStyle("note", base, fontSize=7, leading=9, textColor=MUTED),
    }


def _pdf_text(value) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, (date, datetime)):
        return _fmt_date(value)
    if isinstance(value, (list, tuple)):
        items = [escape(str(v)) for v in value if v]
        if not items:
            return "—"
        return items[0] if len(items) == 1 else "<br/>".join(f"• {v}" for v in items)
    text = str(value)
    if len(text) > _PDF_MAX_CELL_CHARS:
        text = text[:_PDF_MAX_CELL_CHARS].rstrip() + " …"
    return escape(text).replace("\n", "<br/>")


class _NumberedCanvas(rl_canvas.Canvas):
    """Defers page output so every footer can show "page X / total"."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            width, _ = self._pagesize
            self.setFont(_FONT, 7.5)
            self.setFillColor(MUTED)
            self.drawRightString(width - _PDF_MARGIN, 7 * mm, f"Səhifə {self._pageNumber} / {total}")
            super().showPage()
        super().save()


def _pdf_table(cols: list[_Col], rows: list[dict], width: float, st: dict) -> Table:
    num_width = 8 * mm
    total_weight = sum(c.pdf_width for c in cols)
    col_widths = [num_width] + [(width - num_width) * c.pdf_width / total_weight for c in cols]

    data = [[Paragraph("№", st["th"])] + [Paragraph(escape(c.header), st["th"]) for c in cols]]
    for idx, row in enumerate(rows, start=1):
        data.append(
            [Paragraph(str(idx), st["td_center"])]
            + [Paragraph(_pdf_text(c.value(row)), st["td_center" if c.center else "td"]) for c in cols]
        )

    table = Table(data, colWidths=col_widths, repeatRows=1, splitInRow=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
        ("VALIGN", (0, 1), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ZEBRA]),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDER),
        ("LINEBELOW", (0, 0), (-1, 0), 1, NAVY),
        ("BOX", (0, 0), (-1, -1), 0.8, NAVY),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 3.5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3.5),
    ]))
    return table


def render_pdf(doc: ExportDocument) -> bytes:
    _register_fonts()
    st = _styles()
    page_w, page_h = landscape(A4)
    width = page_w - 2 * _PDF_MARGIN
    generated = doc.generated_at.strftime("%d.%m.%Y %H:%M")

    story = []

    # Letterhead: logo on the left, university / title / system on the right,
    # closed by a double navy rule.
    logo_w = 52 * mm
    logo = Image(str(LOGO_PATH), width=logo_w, height=logo_w * _LOGO_H / _LOGO_W)
    title_block = [
        Paragraph(escape(UNIVERSITY_NAME_UPPER), st["kicker"]),
        Spacer(1, 2),
        Paragraph(escape(doc.title), st["title"]),
        Spacer(1, 2),
        Paragraph(escape(SYSTEM_NAME), st["system"]),
    ]
    letterhead = Table([[logo, title_block]], colWidths=[logo_w + 6 * mm, width - logo_w - 6 * mm])
    letterhead.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    story += [
        letterhead,
        Spacer(1, 4),
        HRFlowable(width="100%", thickness=1.8, color=NAVY, spaceBefore=0, spaceAfter=1.5),
        HRFlowable(width="100%", thickness=0.5, color=NAVY, spaceBefore=0, spaceAfter=8),
    ]

    # Details grid: two label/value pairs per row.
    details = list(doc.details) + [("Tərtib tarixi", generated)]
    grid = []
    for i in range(0, len(details), 2):
        pair = details[i:i + 2]
        cells = []
        for label, value in pair:
            cells += [Paragraph(escape(label) + ":", st["label"]), Paragraph(escape(value or "—"), st["value"])]
        cells += [""] * (4 - len(cells))
        grid.append(cells)
    details_table = Table(grid, colWidths=[width * 0.13, width * 0.37, width * 0.13, width * 0.37])
    details_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 1.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
    ]))
    story += [details_table, Spacer(1, 8)]

    # Summary strip.
    summary = _summary(doc)
    stats = Table(
        [[[Paragraph(escape(v), st["stat_value"]), Paragraph(escape(k), st["stat_label"])] for k, v in summary]],
        colWidths=[width / len(summary)] * len(summary),
    )
    stats.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), TINT),
        ("LINEBEFORE", (1, 0), (-1, -1), 1.5, colors.white),
        ("LINEABOVE", (0, 0), (-1, 0), 1.2, NAVY),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story += [stats]

    # Start a new page instead of leaving a heading stranded at the bottom
    # without at least the table header and a row under it.
    heading_room = 32 * mm
    for numeral, heading, _sheet_name, columns, rows, noun in _sections(doc):
        story.append(CondPageBreak(heading_room))
        story.append(Paragraph(f"{numeral}. {escape(heading)}", st["h2"]))
        cols = _visible(columns, doc, "pdf")
        if not rows:
            story.append(Paragraph(escape(EMPTY_TEXT), st["empty"]))
            continue
        if doc.multi_unit:
            for unit, group in groupby(rows, key=lambda r: r.get("unit") or "—"):
                group = list(group)
                story.append(CondPageBreak(heading_room))
                story.append(Paragraph(f"{escape(unit)} — {len(group)} {noun}", st["h3"]))
                story.append(_pdf_table(cols, group, width, st))
        else:
            story.append(_pdf_table(cols, rows, width, st))

    # Signature block.
    signature = Table(
        [[
            Paragraph(escape(doc.signer_label) + ":", st["label"]),
            Paragraph(escape(doc.signer_name or ""), st["value"]),
            Paragraph("İmza: ______________________", st["value"]),
        ]],
        colWidths=[width * 0.12, width * 0.48, width * 0.40],
    )
    signature.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("ALIGN", (2, 0), (2, 0), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(KeepTogether([
        Spacer(1, 18),
        signature,
        Spacer(1, 10),
        Paragraph(
            escape(f"Bu sənəd {SYSTEM_NAME} tərəfindən {generated} tarixində avtomatik formalaşdırılmışdır."),
            st["note"],
        ),
    ]))

    def decorate(canvas, _doc, first_page: bool):
        canvas.saveState()
        canvas.setStrokeColor(NAVY)
        canvas.setLineWidth(0.6)
        canvas.line(_PDF_MARGIN, 11 * mm, page_w - _PDF_MARGIN, 11 * mm)
        canvas.setFont(_FONT, 7.5)
        canvas.setFillColor(MUTED)
        canvas.drawString(_PDF_MARGIN, 7 * mm, f"{UNIVERSITY_NAME} · {SYSTEM_NAME}")
        if not first_page:
            # Running header on continuation pages.
            logo_h = 8 * mm
            canvas.drawImage(
                str(LOGO_PATH), _PDF_MARGIN, page_h - 6 * mm - logo_h,
                width=logo_h * _LOGO_W / _LOGO_H, height=logo_h, mask="auto",
            )
            canvas.setFont(_FONT_BOLD, 8.5)
            canvas.setFillColor(NAVY)
            canvas.drawRightString(page_w - _PDF_MARGIN, page_h - 11 * mm, doc.title)
            canvas.line(_PDF_MARGIN, page_h - 15.5 * mm, page_w - _PDF_MARGIN, page_h - 15.5 * mm)
        canvas.restoreState()

    buf = io.BytesIO()
    template = SimpleDocTemplate(
        buf,
        pagesize=(page_w, page_h),
        leftMargin=_PDF_MARGIN,
        rightMargin=_PDF_MARGIN,
        topMargin=19 * mm,
        bottomMargin=15 * mm,
        title=doc.title,
        author=UNIVERSITY_NAME,
        subject=SYSTEM_NAME,
        creator=SYSTEM_NAME,
    )
    template.build(
        story,
        onFirstPage=lambda c, d: decorate(c, d, True),
        onLaterPages=lambda c, d: decorate(c, d, False),
        canvasmaker=_NumberedCanvas,
    )
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_THIN = Side(style="thin", color=BORDER_HEX)
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor=NAVY_HEX)
_ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA_HEX)


_XLSX_MAX_CELL_CHARS = 32767  # Excel's hard per-cell limit


def _xlsx_value(value):
    if isinstance(value, (list, tuple)):
        value = "\n".join(str(v) for v in value if v) or None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        return value[:_XLSX_MAX_CELL_CHARS]
    return value


def _xlsx_sheet(wb: Workbook, sheet_name: str, columns: list[_Col], rows: list[dict], doc: ExportDocument) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    cols = _visible(columns, doc, "xlsx")
    ncols = len(cols) + 1
    last_letter = get_column_letter(ncols)

    ws.column_dimensions["A"].width = 6
    for i, col in enumerate(cols, start=2):
        ws.column_dimensions[get_column_letter(i)].width = col.xlsx_width

    def full_row(row_idx: int, value, font: Font, align: str = "left", height: float | None = None):
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=ncols)
        cell = ws.cell(row=row_idx, column=1, value=value)
        cell.font = font
        cell.alignment = Alignment(horizontal=align, vertical="center")
        if height:
            ws.row_dimensions[row_idx].height = height

    # Letterhead: the logo floats over the left of rows 1-4, text is right-aligned.
    full_row(1, UNIVERSITY_NAME_UPPER, Font(bold=True, size=13, color=NAVY_HEX), "right", 21)
    full_row(2, doc.title, Font(bold=True, size=12, color=NAVY_HEX), "right", 19)
    full_row(3, SYSTEM_NAME, Font(size=10, color="6B7280"), "right", 17)
    full_row(4, f"Tərtib tarixi: {doc.generated_at.strftime('%d.%m.%Y %H:%M')}", Font(size=9, color="6B7280"), "right", 17)
    logo = XLImage(str(LOGO_PATH))
    logo.height = 70
    logo.width = round(70 * _LOGO_W / _LOGO_H)
    ws.add_image(logo, "A1")
    ws.row_dimensions[5].height = 6
    for c in range(1, ncols + 1):
        ws.cell(row=5, column=c).border = Border(bottom=Side(style="medium", color=NAVY_HEX))

    r = 7
    label_font = InlineFont(b=True, color=NAVY_HEX)
    details = list(doc.details)
    if doc.signer_name:
        details.append((doc.signer_label, doc.signer_name))
    details.append(("Yekun", " · ".join(f"{k}: {v}" for k, v in _summary(doc))))
    for label, value in details:
        full_row(r, CellRichText(TextBlock(label_font, f"{label}: "), value or "—"), Font(size=10))
        r += 1

    r += 1
    header_row = r
    headers = ["№"] + [c.header for c in cols]
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=c, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.border = _CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 32

    if not rows:
        r += 1
        full_row(r, EMPTY_TEXT, Font(italic=True, color="6B7280"))
    for idx, row in enumerate(rows, start=1):
        r += 1
        values = [idx] + [_xlsx_value(c.value(row)) for c in cols]
        centers = [True] + [c.center for c in cols]
        for c, (value, center) in enumerate(zip(values, centers), start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str):
                # openpyxl stores text starting with "=" as a formula; user
                # input must never execute in the reader's Excel.
                cell.data_type = "s"
            cell.border = _CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if center else "left", vertical="top", wrap_text=True)
            if isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
            if idx % 2 == 0:
                cell.fill = _ZEBRA_FILL

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{last_letter}{r}"

    # Print-ready: landscape A4, fit to width, repeat the header row.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.left.text = f"{UNIVERSITY_NAME} · {SYSTEM_NAME}"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Səhifə &P / &N"
    ws.oddFooter.right.size = 8


def render_xlsx(doc: ExportDocument) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for _numeral, _heading, sheet_name, columns, rows, _noun in _sections(doc):
        _xlsx_sheet(wb, sheet_name, columns, rows, doc)
    wb.properties.title = doc.title
    wb.properties.creator = UNIVERSITY_NAME
    wb.properties.subject = SYSTEM_NAME
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
